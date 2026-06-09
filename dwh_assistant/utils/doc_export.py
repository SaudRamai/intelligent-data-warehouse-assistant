"""
doc_export.py — Enterprise Consulting Document Exporters
Produces McKinsey/Deloitte-style PDF and DOCX outputs from structured JSON.
"""

from __future__ import annotations
import io
from datetime import date
from typing import Any

# ---------------------------------------------------------------------------
# PDF EXPORTER — ReportLab Platypus
# ---------------------------------------------------------------------------

from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT, TA_JUSTIFY
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak, KeepTogether
)
from reportlab.platypus.flowables import HRFlowable


# Brand palette (navy + slate)
NAVY        = colors.HexColor("#002244")
SLATE       = colors.HexColor("#334E68")
LIGHT_BLUE  = colors.HexColor("#D5E8F0")
MID_GRAY    = colors.HexColor("#64748B")
LIGHT_GRAY  = colors.HexColor("#F1F5F9")
WHITE       = colors.white
BLACK       = colors.black
ACCENT      = colors.HexColor("#0066CC")


def _build_pdf_styles() -> dict:
    base = getSampleStyleSheet()
    s = {}

    s["cover_title"] = ParagraphStyle(
        "cover_title", fontName="Helvetica-Bold", fontSize=28,
        textColor=WHITE, leading=34, spaceAfter=12, alignment=TA_LEFT
    )
    s["cover_subtitle"] = ParagraphStyle(
        "cover_subtitle", fontName="Helvetica", fontSize=13,
        textColor=colors.HexColor("#B0C4D8"), leading=18, spaceAfter=6, alignment=TA_LEFT
    )
    s["cover_meta"] = ParagraphStyle(
        "cover_meta", fontName="Helvetica", fontSize=10,
        textColor=colors.HexColor("#B0C4D8"), leading=14, alignment=TA_LEFT
    )
    s["section_heading"] = ParagraphStyle(
        "section_heading", fontName="Helvetica-Bold", fontSize=13,
        textColor=NAVY, leading=18, spaceBefore=18, spaceAfter=6,
        borderPad=0
    )
    s["body"] = ParagraphStyle(
        "body", fontName="Helvetica", fontSize=10,
        textColor=colors.HexColor("#1E293B"), leading=16,
        spaceAfter=6, alignment=TA_JUSTIFY
    )
    s["bullet"] = ParagraphStyle(
        "bullet", fontName="Helvetica", fontSize=10,
        textColor=colors.HexColor("#1E293B"), leading=15,
        leftIndent=18, bulletIndent=6, spaceAfter=3
    )
    s["table_header"] = ParagraphStyle(
        "table_header", fontName="Helvetica-Bold", fontSize=9,
        textColor=WHITE, leading=12, alignment=TA_LEFT, wordWrap='CJK'
    )
    s["table_cell"] = ParagraphStyle(
        "table_cell", fontName="Helvetica", fontSize=9,
        textColor=colors.HexColor("#1E293B"), leading=13, alignment=TA_LEFT, wordWrap='CJK'
    )
    s["footer_text"] = ParagraphStyle(
        "footer_text", fontName="Helvetica", fontSize=8,
        textColor=MID_GRAY, alignment=TA_CENTER
    )
    return s


def _parse_content_blocks(content: str, styles: dict) -> list:
    """Parse markdown-like content string into ReportLab flowables."""
    flowables = []
    if not content:
        return flowables

    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            flowables.append(Spacer(1, 4))
        elif stripped.startswith("- ") or stripped.startswith("• "):
            text = stripped[2:].strip()
            # Escape XML special chars
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            flowables.append(
                Paragraph(f'<bullet bulletIndent="6" bulletFontSize="10">\u2022</bullet>{text}',
                           styles["bullet"])
            )
        elif stripped.startswith("**") and stripped.endswith("**"):
            text = stripped[2:-2]
            text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            flowables.append(Paragraph(f"<b>{text}</b>", styles["body"]))
        else:
            text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            # inline bold
            import re
            text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
            flowables.append(Paragraph(text, styles["body"]))

    return flowables


def _build_table_flowable(headers: list, rows: list, styles: dict, page_width: float) -> Table:
    col_count = max(len(headers), max((len(r) for r in rows), default=0)) if (headers or rows) else 1
    
    base_font_size = 9
    base_leading = 13
    padding = 8
    
    if col_count >= 8:
        base_font_size = 6
        base_leading = 8
        padding = 3
    elif col_count >= 6:
        base_font_size = 7
        base_leading = 9
        padding = 4
        
    from copy import deepcopy
    header_style = deepcopy(styles["table_header"])
    cell_style = deepcopy(styles["table_cell"])
    
    header_style.fontSize = base_font_size + 1
    header_style.leading = base_leading + 1
    cell_style.fontSize = base_font_size
    cell_style.leading = base_leading

    header_row = [Paragraph(str(h), header_style) for h in headers]
    data = [header_row]
    for row in rows:
        data.append([Paragraph(str(c), cell_style) for c in row])

    # Calculate proportional column widths based on content length
    max_lens = [0] * col_count
    for i, h in enumerate(headers):
        max_lens[i] = max(max_lens[i], len(str(h)))
    for r in rows:
        for i, c in enumerate(r):
            if i < col_count:
                max_lens[i] = max(max_lens[i], len(str(c)))
                
    max_lens = [max(l, 5) for l in max_lens]
    total_len = sum(max_lens)
    usable_width = page_width - 2 * inch
    
    col_widths = [usable_width * (l / total_len) for l in max_lens]
    
    # Enforce minimum width
    min_width = 30
    for i in range(len(col_widths)):
        if col_widths[i] < min_width:
            col_widths[i] = min_width
            
    # Normalize back to usable width
    total_width = sum(col_widths)
    if total_width > usable_width:
        scale = usable_width / total_width
        col_widths = [w * scale for w in col_widths]

    tbl = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0),  NAVY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LIGHT_GRAY]),
        ("GRID",          (0, 0), (-1, -1),  0.4, colors.HexColor("#CBD5E1")),
        ("TOPPADDING",    (0, 0), (-1, -1),  padding),
        ("BOTTOMPADDING", (0, 0), (-1, -1),  padding),
        ("LEFTPADDING",   (0, 0), (-1, -1),  padding),
        ("RIGHTPADDING",  (0, 0), (-1, -1),  padding),
        ("VALIGN",        (0, 0), (-1, -1),  "TOP"),
        ("LINEBELOW",     (0, 0), (-1, 0),   1, ACCENT),
    ]))
    return tbl


class _PageDecorator:
    """Adds header/footer on every page."""

    def __init__(self, title: str, client: str, doc_date: str):
        self.title = title
        self.client = client
        self.doc_date = doc_date

    def __call__(self, canvas, doc):
        canvas.saveState()
        w, h = letter

        # Header bar
        canvas.setFillColor(NAVY)
        canvas.rect(0, h - 44, w, 44, fill=1, stroke=0)
        canvas.setFillColor(WHITE)
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawString(0.5 * inch, h - 28, self.title)
        canvas.setFont("Helvetica", 9)
        canvas.drawRightString(w - 0.5 * inch, h - 28, f"Prepared for {self.client}")

        # Footer
        canvas.setFillColor(LIGHT_GRAY)
        canvas.rect(0, 0, w, 28, fill=1, stroke=0)
        canvas.setFillColor(MID_GRAY)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(0.5 * inch, 9, f"Confidential — {self.doc_date}")
        canvas.drawCentredString(w / 2, 9, f"Page {doc.page}")
        canvas.drawRightString(w - 0.5 * inch, 9, "© All Rights Reserved")

        canvas.restoreState()


class ConsultingPDFExporter:
    """Generates a McKinsey-style PDF from structured section JSON."""

    def __init__(self, metadata: dict):
        self.title = metadata.get("title", "Solution Document")
        self.client = metadata.get("client", "Client")
        self.version = metadata.get("version", "1.0")
        self.doc_date = metadata.get("date", date.today().isoformat())

    def generate(self, sections: list[dict]) -> bytes:
        buf = io.BytesIO()
        w, h = letter
        decorator = _PageDecorator(self.title, self.client, self.doc_date)

        doc = SimpleDocTemplate(
            buf, pagesize=letter,
            leftMargin=0.75 * inch, rightMargin=0.75 * inch,
            topMargin=0.85 * inch, bottomMargin=0.65 * inch,
            title=self.title, author="Data Consulting Practice"
        )

        styles = _build_pdf_styles()
        story = []

        # ── Cover page ───────────────────────────────────────────────────────
        story.append(Spacer(1, 1.2 * inch))

        # Navy cover block rendered as a table (background trick)
        cover_data = [[
            Paragraph(self.title, styles["cover_title"]),
        ]]
        cover_tbl = Table(cover_data, colWidths=[w - 1.5 * inch])
        cover_tbl.setStyle(TableStyle([
            ("BACKGROUND",   (0, 0), (-1, -1), NAVY),
            ("TOPPADDING",   (0, 0), (-1, -1), 28),
            ("BOTTOMPADDING",(0, 0), (-1, -1), 28),
            ("LEFTPADDING",  (0, 0), (-1, -1), 24),
            ("RIGHTPADDING", (0, 0), (-1, -1), 24),
        ]))
        story.append(cover_tbl)
        story.append(Spacer(1, 14))

        meta_lines = [
            f"<b>Prepared for:</b> {self.client}",
            f"<b>Version:</b> {self.version}",
            f"<b>Date:</b> {self.doc_date}",
            "<b>Classification:</b> Confidential",
        ]
        for m in meta_lines:
            story.append(Paragraph(m, styles["body"]))
        story.append(HRFlowable(width="100%", thickness=2, color=NAVY, spaceAfter=4))
        story.append(PageBreak())

        # ── Sections ─────────────────────────────────────────────────────────
        for sec in sections:
            heading = sec.get("heading", "")
            sec_type = sec.get("type", "text")

            block = []

            if heading:
                block.append(Paragraph(heading, styles["section_heading"]))
                block.append(HRFlowable(
                    width="100%", thickness=1,
                    color=ACCENT, spaceAfter=6
                ))

            if sec_type == "table":
                headers = sec.get("headers", [])
                rows = sec.get("rows", [])
                if headers or rows:
                    block.append(Spacer(1, 4))
                    block.append(_build_table_flowable(headers, rows, styles, w))
                    block.append(Spacer(1, 8))
            else:
                content = sec.get("content", "")
                block.extend(_parse_content_blocks(content, styles))

            story.append(KeepTogether(block[:4]))  # keep heading + first lines together
            story.extend(block[4:] if len(block) > 4 else [])
            story.append(Spacer(1, 10))

        doc.build(story, onFirstPage=decorator, onLaterPages=decorator)
        return buf.getvalue()


# ---------------------------------------------------------------------------
# DOCX EXPORTER — python-docx
# ---------------------------------------------------------------------------

from docx import Document as DocxDocument
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy


_NAVY_RGB   = RGBColor(0x00, 0x22, 0x44)
_ACCENT_RGB = RGBColor(0x00, 0x66, 0xCC)
_GRAY_RGB   = RGBColor(0x64, 0x74, 0x8B)
_LIGHT_RGB  = RGBColor(0xF1, 0xF5, 0xF9)


def _hex_to_docx_color(hex_str: str) -> str:
    """Convert #RRGGBB → 'RRGGBB' for OxmlElement shading."""
    return hex_str.lstrip("#")


def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, color="CBD5E1"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _add_heading_paragraph(doc: DocxDocument, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    run = p.runs[0] if p.runs else p.add_run(text)
    run.font.color.rgb = _NAVY_RGB
    run.font.name = "Arial"
    # Underline rule after heading
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0066CC")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _add_body_paragraph(doc: DocxDocument, text: str, bold: bool = False, bullet: bool = False):
    if bullet:
        p = doc.add_paragraph(style="List Bullet")
    else:
        p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Arial"
    if bold:
        run.bold = True
    p.paragraph_format.space_after = Pt(4)
    return p


def _parse_and_add_content(doc: DocxDocument, content: str):
    import re
    for line in content.splitlines():
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph()
            continue
        if stripped.startswith("- ") or stripped.startswith("• "):
            _add_body_paragraph(doc, stripped[2:].strip(), bullet=True)
        elif stripped.startswith("**") and stripped.endswith("**"):
            _add_body_paragraph(doc, stripped[2:-2], bold=True)
        else:
            p = doc.add_paragraph()
            # handle inline **bold**
            parts = re.split(r'\*\*(.+?)\*\*', stripped)
            for i, part in enumerate(parts):
                run = p.add_run(part)
                run.font.size = Pt(10)
                run.font.name = "Arial"
                if i % 2 == 1:
                    run.bold = True
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _add_table(doc: DocxDocument, headers: list, rows: list):
    col_count = max(len(headers), max((len(r) for r in rows), default=0)) if (headers or rows) else 1
    total_rows = 1 + len(rows)

    tbl = doc.add_table(rows=total_rows, cols=col_count)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Auto-width
    tbl.autofit = True

    # Header row
    hdr_row = tbl.rows[0]
    for idx, h in enumerate(headers):
        cell = hdr_row.cells[idx]
        _set_cell_bg(cell, "002244")
        p = cell.paragraphs[0]
        run = p.add_run(str(h))
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = "Arial"

    # Data rows
    for r_idx, row in enumerate(rows):
        bg = "F1F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_val in enumerate(row):
            if c_idx >= col_count:
                break
            cell = tbl.rows[r_idx + 1].cells[c_idx]
            _set_cell_bg(cell, bg)
            _set_cell_borders(cell)
            p = cell.paragraphs[0]
            run = p.add_run(str(cell_val))
            run.font.size = Pt(9)
            run.font.name = "Arial"

    doc.add_paragraph()  # spacer after table


def _add_cover_page(doc: DocxDocument, title: str, client: str, version: str, doc_date: str):
    # Large navy title block
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(72)
    p.paragraph_format.space_after = Pt(6)
    # Blue left border accent
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "12")
    left.set(qn("w:color"), "002244")
    pBdr.append(left)
    pPr.append(pBdr)

    run = p.add_run(title)
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.name = "Arial"
    run.font.color.rgb = _NAVY_RGB

    doc.add_paragraph()

    for label, value in [
        ("Prepared for", client),
        ("Version",      version),
        ("Date",         doc_date),
        ("Classification", "Confidential"),
    ]:
        mp = doc.add_paragraph()
        r1 = mp.add_run(f"{label}: ")
        r1.bold = True
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = _GRAY_RGB
        r2 = mp.add_run(value)
        r2.font.name = "Arial"
        r2.font.size = Pt(10)

    doc.add_page_break()


class ConsultingDOCXExporter:
    """Generates an enterprise-grade DOCX from structured section JSON."""

    def __init__(self, metadata: dict):
        self.title = metadata.get("title", "Solution Document")
        self.client = metadata.get("client", "Client")
        self.version = metadata.get("version", "1.0")
        self.doc_date = metadata.get("date", date.today().isoformat())

    def generate(self, sections: list[dict]) -> bytes:
        doc = DocxDocument()

        # Page margins
        for section in doc.sections:
            section.top_margin    = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin   = Inches(1.0)
            section.right_margin  = Inches(1.0)

        # Default font
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)

        _add_cover_page(doc, self.title, self.client, self.version, self.doc_date)

        for sec in sections:
            heading = sec.get("heading", "")
            sec_type = sec.get("type", "text")

            if heading:
                _add_heading_paragraph(doc, heading, level=2)

            if sec_type == "table":
                headers = sec.get("headers", [])
                rows    = sec.get("rows", [])
                if headers or rows:
                    _add_table(doc, headers, rows)
            else:
                content = sec.get("content", "")
                _parse_and_add_content(doc, content)

            doc.add_paragraph()  # section spacer

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
